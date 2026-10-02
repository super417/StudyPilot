"""Document parsing, text chunking, and storage (requirements 11.5-11.9)."""

from __future__ import annotations

import io
import uuid

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.entities import UserDocument

MAX_FILE_BYTES = 20 * 1024 * 1024
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
ALLOWED_FILE_TYPES = frozenset({"pdf", "docx", "txt", "md"})

# Upper bound on the retrieved document context assembled for planning, so an
# unusually large document set cannot bloat the planning prompt (requirement
# 9.8/17.4). Retrieval truncates to this many characters when a caller does not
# override it.
MAX_CONTEXT_CHARS = 20000


class UnsupportedFileTypeError(ValueError):
    """Raised when an uploaded file's extension is not a supported type."""

    code = "UNSUPPORTED_TYPE"

    def __init__(self, message: str = "unsupported file type") -> None:
        super().__init__(message)


class DocumentParseError(ValueError):
    """Raised when a document cannot be parsed or yields no usable text."""

    code = "PARSE_FAILED"

    def __init__(self, message: str = "document parsing failed") -> None:
        super().__init__(message)


def detect_file_type(filename: str) -> str:
    """Return the supported file type derived from a filename extension.

    Only the lowercased extension is inspected. Unknown, missing, or
    unsupported extensions raise :class:`UnsupportedFileTypeError`.
    """
    _, _, extension = (filename or "").rpartition(".")
    normalized = extension.strip().lower()
    if not normalized or normalized not in ALLOWED_FILE_TYPES:
        raise UnsupportedFileTypeError()
    return normalized


def _decode_text(data: bytes) -> str:
    """Decode UTF-8 text, dropping a leading BOM, else fail safely."""
    try:
        # utf-8-sig strips a leading BOM if present and otherwise behaves
        # exactly like utf-8.
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentParseError() from error


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def parse_document(file_type: str, data: bytes) -> str:
    """Extract plain text from ``data`` according to ``file_type``.

    Any underlying parser failure is translated into a generic
    :class:`DocumentParseError` so that low-level details never leak.
    """
    try:
        if file_type in ("txt", "md"):
            text = _decode_text(data)
        elif file_type == "pdf":
            text = _parse_pdf(data)
        elif file_type == "docx":
            text = _parse_docx(data)
        else:  # pragma: no cover - guarded by detect_file_type upstream.
            raise UnsupportedFileTypeError()
    except DocumentParseError:
        raise
    except Exception as error:  # noqa: BLE001 - normalize to a safe error.
        raise DocumentParseError() from error
    return text.strip()


def chunk_text(
    text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP
) -> list[str]:
    """Split ``text`` into overlapping chunks (Property 15).

    Each chunk holds at most ``size`` characters; adjacent chunks overlap by
    exactly ``overlap`` characters (the final chunk may be shorter). Removing
    the overlap and concatenating in order reproduces the original text with no
    gaps. Empty or whitespace-only input yields an empty list.
    """
    if not text or not text.strip():
        return []

    step = size - overlap
    chunks: list[str] = []
    start = 0
    length = len(text)
    while start < length:
        chunks.append(text[start : start + size])
        if start + size >= length:
            break
        start += step
    return chunks


def list_user_documents(session: Session, user_id: uuid.UUID) -> list[dict]:
    """Aggregate chunk rows into per-document summaries for the owning user.

    Each document is ready when it has at least one chunk. Ordered by the
    newest ``created_at`` among its chunks (most recent upload first).
    """
    rows = list(
        session.scalars(
            select(UserDocument)
            .where(UserDocument.user_id == user_id)
            .order_by(UserDocument.created_at.desc(), UserDocument.chunk_index.asc())
        )
    )
    by_doc: dict[str, dict] = {}
    for row in rows:
        existing = by_doc.get(row.doc_id)
        if existing is None:
            by_doc[row.doc_id] = {
                "docId": row.doc_id,
                "filename": row.filename,
                "fileType": row.file_type,
                "chunks": 1,
                "ready": True,
                "uploadedAt": row.created_at.isoformat(),
            }
        else:
            existing["chunks"] = int(existing["chunks"]) + 1
            # Keep uploadedAt from the newest-first first sighting.
    return list(by_doc.values())


def store_document(
    session: Session, user_id: uuid.UUID, filename: str, data: bytes
) -> tuple[str, int]:
    """Parse, chunk, and atomically persist a document; return (doc_id, count).

    Nothing is written when the type is unsupported, the file is too large, or
    parsing yields empty text (Property 16 write atomicity).
    """
    file_type = detect_file_type(filename)

    if len(data) > MAX_FILE_BYTES:
        raise DocumentParseError("文件过大")

    text = parse_document(file_type, data)
    chunks = chunk_text(text)
    if not chunks:
        raise DocumentParseError()

    doc_id = uuid.uuid4().hex
    session.add_all(
        UserDocument(
            user_id=user_id,
            doc_id=doc_id,
            filename=filename,
            file_type=file_type,
            chunk_index=index,
            content=chunk,
        )
        for index, chunk in enumerate(chunks)
    )
    try:
        session.commit()
    except Exception:
        session.rollback()
        raise
    return doc_id, len(chunks)


def filter_ready_documents(
    session: Session, user_id: uuid.UUID, document_ids: list[str] | None
) -> tuple[list[str], list[str]]:
    """Split selected doc ids into (ready, skipped) for a user (needs 9.7/17.4).

    A document is *ready* when the user owns at least one ``User_Documents`` row
    for that ``doc_id`` — i.e. parsing and chunking have completed. Unknown ids
    and ids belonging to another user are *skipped*. The input order is
    preserved and duplicates are collapsed to their first occurrence. Empty or
    missing input yields ``([], [])``.
    """
    ready: list[str] = []
    skipped: list[str] = []
    seen: set[str] = set()

    for doc_id in document_ids or []:
        if doc_id in seen:
            continue
        seen.add(doc_id)
        exists = session.scalar(
            select(UserDocument.id)
            .where(
                UserDocument.user_id == user_id,
                UserDocument.doc_id == doc_id,
            )
            .limit(1)
        )
        if exists is not None:
            ready.append(doc_id)
        else:
            skipped.append(doc_id)

    return ready, skipped


def retrieve_document_chunks(
    session: Session,
    user_id: uuid.UUID,
    doc_ids: list[str] | None,
    max_chars: int | None = None,
) -> str:
    """Return the concatenated chunk text for a user's documents as context.

    Chunks are grouped per ``doc_id`` and ordered by ``chunk_index`` ascending,
    then joined into a single plain-text block. Only rows owned by ``user_id``
    contribute, so another user's content can never leak in. Nothing but the
    stored text is returned (no ids, filenames, or metadata). The result is
    truncated to ``max_chars`` characters (defaulting to ``MAX_CONTEXT_CHARS``)
    to keep the planning prompt bounded.
    """
    if not doc_ids:
        return ""

    limit = MAX_CONTEXT_CHARS if max_chars is None else max_chars

    parts: list[str] = []
    for doc_id in doc_ids:
        contents = session.scalars(
            select(UserDocument.content)
            .where(
                UserDocument.user_id == user_id,
                UserDocument.doc_id == doc_id,
            )
            .order_by(UserDocument.chunk_index)
        ).all()
        parts.extend(contents)

    text = "\n".join(parts)
    if limit is not None and limit >= 0:
        text = text[:limit]
    return text
